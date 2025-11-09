from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from django.contrib import messages
from django.db.models import Count
from django.db.models.functions import TruncDate
from django.utils.timezone import make_naive
from openpyxl import Workbook, load_workbook
import io, json

from .models import Employee, Task
from .forms import EmployeeForm, TaskForm
from django.db.models import Q   # ✅ add this


# -------------------------
# Dashboard
# -------------------------
from django.shortcuts import render
from django.db.models import Count
from .models import Employee, Task
from django.utils.timezone import localdate
from django.db.models.functions import TruncDate

from django.db.models import Count
from django.utils import timezone
import json

def dashboard(request):
    # ✅ Recent employees limit 5
    recent_employees = Employee.objects.order_by('-id')[:5]

    # ✅ Task status count
    task_status_data = list(Task.objects.values('status').annotate(count=Count('status')))

    # ✅ Employee vs Tasks
    employee_task_data = list(
        Employee.objects.annotate(task_count=Count('task')).values('name', 'task_count')
    )
    # Extra metrics
    upcoming_tasks = Task.objects.filter(due_date__gte=timezone.now().date()).count()

    overdue_tasks = Task.objects.filter(due_date__lt=timezone.now().date()).exclude(status="Completed").count()



    # ✅ Task Trend last 7 days
    today = timezone.now().date()
    trend = []
    for i in range(7):
        day = today - timezone.timedelta(days=i)
        count = Task.objects.filter(created_at__date=day).count()
        trend.append({"date": day.strftime("%d %b"), "count": count})
        trend.reverse()

    return render(request, 'employees/dashboard.html', {
        "recent_employees": recent_employees,
        "task_status_data": json.dumps(task_status_data),
        "employee_task_data": json.dumps(employee_task_data),
        "task_trend": json.dumps(trend),
        "total_employees": Employee.objects.count(),
        "total_tasks": Task.objects.count(),
        "completed_tasks": Task.objects.filter(status="Completed").count(),
        "upcoming_tasks": upcoming_tasks,
        "overdue_tasks": overdue_tasks,

    })
    


# -------------------------
# Employee CRUD
# -------------------------
def employee_list(request):
    search = request.GET.get("search", "")

    if search:
        employees = Employee.objects.filter(
            Q(name__icontains=search) |
            Q(email__icontains=search) |
            Q(role__icontains=search)
        ).order_by('-id')
    else:
        employees = Employee.objects.all().order_by('-id')

    return render(request, 'employees/employee_list.html', {
        "employees": employees,
        "search": search
    })


def add_employee(request):
    form = EmployeeForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Employee added ✅")
        return redirect('employee_list')
    return render(request, 'employees/add_employee.html', {'form': form})

def edit_employee(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    form = EmployeeForm(request.POST or None, instance=employee)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Employee updated ✅")
        return redirect('employee_list')
    return render(request, 'employees/add_employee.html', {'form': form})

def delete_employee(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    if request.method == "POST":
        employee.delete()
        messages.success(request, "Employee deleted ❌")
        return redirect('employee_list')
    return render(request, 'employees/delete_confirm.html', {'employee': employee})

# -------------------------
# Employee Excel Import
# -------------------------
def import_excel(request):
    if request.method == "POST" and request.FILES.get("file"):
        wb = load_workbook(request.FILES["file"])
        sheet = wb.active
        added = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0] or not row[1]:  # skip empty rows
                continue

            name, email = row[0], row[1]
            role = row[2] if len(row) > 2 else ""

            if not Employee.objects.filter(email=email).exists():
                Employee.objects.create(name=name, email=email, role=role)
                added += 1

        messages.success(request, f"{added} employees imported ✅")
        return redirect('employee_list')

    return render(request, 'employees/import_excel.html')

# -------------------------
# Employee Excel Export
# -------------------------
def export_excel(request):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Employees"
    sheet.append(["ID", "Name", "Email", "Phone", "Joined Date"])

    for emp in Employee.objects.all():
        joined = make_naive(emp.joined_date) if emp.joined_date else ""
        sheet.append([emp.id, emp.name, emp.email, emp.phone, joined])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(buffer,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename=employees.xlsx'
    return response

# -------------------------
# Task CRUD
# -------------------------
def task_list(request):
    search = request.GET.get("search", "")

    if search:
        tasks = Task.objects.filter(
            Q(title__icontains=search) |
            Q(employee__name__icontains=search) |
            Q(status__icontains=search)
        ).order_by('-created_at')
    else:
        tasks = Task.objects.all().order_by('-created_at')

    return render(request, 'employees/task_list.html', {
        "tasks": tasks,
        "search": search
    })


def add_task(request):
    form = TaskForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Task added ✅")
        return redirect('task_list')
    return render(request, 'employees/add_task.html', {'form': form})

def edit_task(request, pk):
    task = get_object_or_404(Task, pk=pk)
    form = TaskForm(request.POST or None, instance=task)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Task updated ✅")
        return redirect('task_list')
    return render(request, 'employees/add_task.html', {'form': form})

def delete_task(request, pk):
    task = get_object_or_404(Task, pk=pk)
    if request.method == "POST":
        task.delete()
        messages.success(request, "Task deleted ❌")
        return redirect('task_list')
    return render(request, 'employees/delete_confirm.html', {'task': task})

# -------------------------
# Task Excel Import
# -------------------------
def import_tasks_excel(request):
    if request.method == "POST" and request.FILES.get("file"):
        wb = load_workbook(request.FILES["file"])
        sheet = wb.active
        added = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0] or not row[1]:
                continue

            employee_email, title, description, due_date, status = row

            try:
                employee = Employee.objects.get(email=employee_email)
            except Employee.DoesNotExist:
                continue

            Task.objects.create(
                employee=employee,
                title=title,
                description=description,
                due_date=due_date,
                status=status
            )
            added += 1

        messages.success(request, f"{added} tasks imported ✅")
        return redirect('task_list')

    return render(request, 'employees/import_excel.html', {"type": "tasks"})

# -------------------------
# Task Excel Export
# -------------------------
def export_tasks_excel(request):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Tasks"
    sheet.append(["ID", "Employee Email", "Title", "Description", "Due Date", "Status", "Created At"])

    for task in Task.objects.all():
        due = make_naive(task.due_date) if task.due_date else ""
        created = make_naive(task.created_at) if task.created_at else ""
        sheet.append([task.id, task.employee.email, task.title, task.description, due, task.status, created])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(buffer,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename=tasks.xlsx'
    return response
